"""Find report fields by their printed labels and merged table boundaries."""
import re
import unicodedata

from openpyxl.cell.cell import MergedCell


def label_key(value):
    text = unicodedata.normalize('NFKC', str(value or '')).lower()
    text = text.replace('ي', 'ی').replace('ك', 'ک')
    return re.sub(r'[^\w]+', '', text)


class LayoutError(ValueError):
    pass


class ReportLayout:
    def __init__(self, sheet):
        self.sheet = sheet
        self.cells = [c for row in sheet for c in row if not isinstance(c, MergedCell)]
        self.merges = list(sheet.merged_cells.ranges)

    def bounds(self, cell):
        for area in self.merges:
            if cell.coordinate in area:
                return area.min_row, area.min_col, area.max_row, area.max_col
        return cell.row, cell.column, cell.row, cell.column

    def matches(self, labels, *, start=None, end=None, column=None, prefix=False):
        if isinstance(labels, str):
            labels = (labels,)
        keys = [label_key(s) for s in labels]
        return [c for c in self.cells if isinstance(c.value, str)
                and (start is None or c.row >= start) and (end is None or c.row < end)
                and (column is None or c.column == column)
                and any(label_key(c.value).startswith(k) if prefix else label_key(c.value) == k for k in keys)]

    def find(self, labels, **scope):
        matches = self.matches(labels, **scope)
        if len(matches) != 1:
            raise LayoutError(f'{self.sheet.title}: expected one label {labels!r}, found {len(matches)}')
        return matches[0]

    def right(self, label, count=1):
        cell = label
        for _ in range(count):
            row, _, _, end = self.bounds(cell)
            cell = self.sheet.cell(row, end + 1)
            if isinstance(cell, MergedCell):
                raise LayoutError(f'{self.sheet.title}: value after {label.value!r} is not a cell anchor')
        return cell

    def below(self, label):
        _, col, end, _ = self.bounds(label)
        cell = self.sheet.cell(end + 1, col)
        if isinstance(cell, MergedCell):
            raise LayoutError(f'{self.sheet.title}: value below {label.value!r} is not a cell anchor')
        return cell

    def value(self, labels, **scope):
        return self.right(self.find(labels, **scope)).value

    def region_rows(self, heading, stop_labels, *, column=None):
        """Rows below a heading, stopping at the next named section in its lane."""
        start = self.bounds(heading)[2] + 1
        stops = self.matches(stop_labels, start=start, column=column)
        if not stops:
            raise LayoutError(f'{self.sheet.title}: section end missing after {heading.value!r}')
        return range(start, min(c.row for c in stops))

    def column(self, heading, labels, *, before=None):
        """Resolve a header in this section, without relying on its column number."""
        return self.find(labels, start=heading.row, end=before).column

    def slots(self, heading, *, count=None):
        """Existing value anchors in the lane immediately to the right of a section label."""
        start, _, end, _ = self.bounds(heading)
        col = self.right(heading).column
        result = [self.sheet.cell(r, col) for r in range(start, end + 1)
                  if not isinstance(self.sheet.cell(r, col), MergedCell)]
        return result[:count] if count else result
