import re
import json
import shutil
import tempfile
import unicodedata
from collections import Counter, defaultdict
from collections.abc import Callable
from pathlib import Path
from typing import Any

import openpyxl
from pypdf import PdfReader

from my_qt_app.workbook_writer import check_saved_values, save_report_package
from my_qt_app.excel_sources import read_company, read_pob, read_service
from my_qt_app.target_report import TargetReport


def normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    return " ".join(text.replace("ي", "ی").replace("ك", "ک").split()).strip(" :").lower()


def number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = normalize(value).replace(",", "").replace("٫", ".")
    if text == "صفر":
        return 0.0
    if re.fullmatch(r"\d+(?:\.\d+)?", text):
        return float(text)
    return None


class Consolidator:
    def __init__(
        self,
        folder_path: Path,
        log_callback: Callable[[str], None] | None = None,
        status_callback: Callable[[str], None] | None = None,
    ):
        self.folder_path = Path(folder_path)
        self.log_callback = log_callback
        self.status_callback = status_callback
        self.outputs: list[Path] = []

    def _log(self, logs: list[str], message: str) -> None:
        logs.append(message)
        if self.log_callback:
            self.log_callback(message)

    def _set_status(self, status: str) -> None:
        if self.status_callback:
            self.status_callback(status)

    def extract_text(self, path: Path, *, layout: bool = False) -> str:
        if path.suffix.lower() == ".pdf":
            kwargs = {"extraction_mode": "layout", "layout_mode_space_vertically": False} if layout else {}
            text = "\n".join(page.extract_text(**kwargs) or "" for page in PdfReader(path).pages)
            if not text.strip():
                raise ValueError(f"No readable text in {path.name}; a scanned PDF needs OCR.")
            return unicodedata.normalize("NFKC", text)
        raise ValueError(f"Expected a PDF: {path.name}")

    def locate_files(self) -> dict[str, Path | None]:
        candidates: dict[str, list[Path]] = {k: [] for k in (
            "target", "company_ddr", "pob", "mdr", "edr", "mud_report"
        )}
        for path in sorted(self.folder_path.iterdir()):
            name = path.stem.lower()
            if not path.is_file() or path.suffix.lower() not in {".xlsx", ".pdf"}:
                continue
            if name.startswith("~$") or "consolidated" in name:
                continue
            if re.search(r"tool[ _-]*(?:poosher|pusher)", name):
                key = "target"
            elif "pob" in name:
                key = "pob"
            elif re.search(r"mud[ _-]*report", name):
                key = "mud_report"
            elif "mdr" in name:
                key = "mdr"
            elif "edr" in name:
                key = "edr"
            elif "ddr" in name:
                key = "company_ddr"
            else:
                continue
            if key in {"target", "company_ddr", "pob"} and path.suffix.lower() != ".xlsx":
                continue
            candidates[key].append(path)

        files: dict[str, Path | None] = {}
        for key, paths in candidates.items():
            # Revisions take priority; within a revision use the structured workbook.
            def rank(path: Path) -> tuple[int, bool]:
                revision = re.search(r"rev[ _.-]*(\d+)", path.stem, re.IGNORECASE)
                return (int(revision[1]) if revision else 0, path.suffix.lower() == ".xlsx")
            best = max(map(rank, paths), default=None)
            winners = [p for p in paths if rank(p) == best]
            if len(winners) > 1:
                raise ValueError(f"Ambiguous {key} reports: {', '.join(p.name for p in winners)}")
            files[key] = winners[0] if winners else None
        return files

    def extract_pob(self, path: Path) -> dict[str, Any]:
        return read_pob(path)

    def parse_company_man_ddr(self, path: Path) -> dict[str, Any]:
        return read_company(path)

    def _pdf_hours(self, text: str) -> dict[str, float]:
        """Match each runtime to its own printed heading, including RTL reports."""
        lines = text.splitlines()
        result = {}
        for index, line in enumerate(lines[:-1]):
            headings = []
            for chunk in re.finditer(r"\S(?:.*?\S)?(?=\s{2,}|$)", line):
                label = normalize(chunk[0])
                n = re.search(r"[1-4]", label)
                if "پمک روتوم" in label and n:
                    key = f"camp_{n[0]}"
                elif "روتوم" in label and n:
                    key = f"gen_{n[0]}"
                elif "روسرپمک" in label and n:
                    key = f"comp_{n[0]}"
                elif label == "یرارطضا":
                    key = "emergency"
                else:
                    continue
                headings.append((key, (chunk.start() + chunk.end()) / 2))
            if len(headings) < 3:
                continue
            values = list(re.finditer(r"\S+", lines[index + 1]))
            if len(values) != len(headings) or any(number(v[0]) is None for v in values):
                continue
            assignments = defaultdict(list)
            for value in values:
                center = (value.start() + value.end()) / 2
                key = min(headings, key=lambda h: abs(h[1] - center))[0]
                assignments[key].append(number(value[0]))
            if all(len(assignments[key]) == 1 for key, _ in headings):
                result.update({key: assignments[key][0] for key, _ in headings})
        return result

    def _pdf_remarks(self, path: Path, *, mud: bool = False) -> list[str]:
        remarks = []
        for page in PdfReader(path).pages:
            fragments = []
            def visit(text, cm, tm, font, size):
                if text.strip():
                    fragments.append((tm[4], tm[5], unicodedata.normalize("NFKC", text).strip()))
            page.extract_text(visitor_text=visit)
            upper = [y for x, y, t in fragments if ("کارهای" if mud else "شرح اقدامات") in t]
            lower = [y for x, y, t in fragments if "قطعات مصرف" in t]
            if not upper or (not mud and not lower):
                continue
            top, bottom = max(upper), max(lower) if lower and not mud else 0
            rows = defaultdict(list)
            for x, y, text in fragments:
                if bottom < y < top and (not mud or x > float(page.mediabox.width) / 2):
                    rows[round(y)].append((x, text))
            numbered_paragraphs = False
            for y in sorted(rows, reverse=True):
                # Keep number strings intact while ordering Persian fragments right to left.
                line = " ".join(t for x, t in sorted(rows[y], reverse=True))
                line = re.sub(r"^\s*\d+\s*[-_]\s*", "", line).strip()
                if line and not line.startswith("گزارش واحد"):
                    numbered = re.fullmatch(r"\.\s*(.*?)\s+(\d+)", line)
                    if numbered and not mud:
                        remarks.append(numbered[1])
                        numbered_paragraphs = True
                    elif not mud and remarks and numbered_paragraphs:
                        # Mechanical PDFs wrap each numbered item across lines.
                        remarks[-1] += "\n" + line
                    else:
                        remarks.append(line)
        return remarks

    def parse_electrical_report(self, path: Path) -> dict[str, Any]:
        if path.suffix.lower() == ".xlsx":
            return self._parse_service_workbook(path)
        text = self.extract_text(path, layout=True)
        date = re.search(r"\d{4}/\d{2}/\d{2}", text)
        location = next((line for line in text.splitlines() if "هاگتسد لحم" in normalize(line) or "محل دستگاه" in normalize(line)), "")
        well = re.search(r"S[ _-]*(\d+)", location, re.I)
        return {"remarks": self._pdf_remarks(path), "runtimes": self._pdf_hours(text),
                "date": date[0] if date else None, "well_number": well[1] if well else None}

    def _parse_service_workbook(self, path: Path) -> dict[str, Any]:
        return read_service(path)

    def parse_mechanical_report(self, path: Path) -> dict[str, Any]:
        if path.suffix.lower() == ".xlsx":
            return self._parse_service_workbook(path)
        return self.parse_electrical_report(path)

    def parse_mud_report(self, path: Path) -> dict[str, Any]:
        if path.suffix.lower() != ".pdf":
            raise ValueError("Mud service workbook layout is not mapped; supply its PDF.")
        text = self.extract_text(path, layout=True)
        runtimes, screens, equipment = {}, {}, {}
        shaker_columns = []
        decks = defaultdict(list)
        for line in text.splitlines():
            chunks = list(re.finditer(r"\S(?:.*?\S)?(?=\s{2,}|$)", line))
            if line.lstrip().startswith("SHALE SHAKERS"):
                for chunk in chunks:
                    label = re.fullmatch(r"Shaker\s+([1-3])", chunk[0])
                    if label:
                        shaker_columns.append((label[1], (chunk.start() + chunk.end()) / 2))
                    elif chunk[0] == "Mud Cleaner":
                        shaker_columns.append(("cleaner_1", (chunk.start() + chunk.end()) / 2))
            if line.lstrip().startswith("Screens Size") and shaker_columns:
                centers = [position for _, position in shaker_columns]
                right = centers[-1] + (centers[-1] - centers[-2]) / 2
                for chunk in chunks[1:]:
                    center = (chunk.start() + chunk.end()) / 2
                    if center < right and number(chunk[0]) is not None:
                        key = min(shaker_columns, key=lambda item: abs(item[1] - center))[0]
                        decks[key].append(chunk[0])
            for heading, key in (("SENIOR MUD EN", "senior_mud_engineer"), ("JUNIOR MUD EN", "junior_mud_engineer")):
                if line.lstrip().startswith(heading):
                    columns = re.split(r"\s{2,}", line.strip())
                    if len(columns) >= 2:
                        equipment[key] = columns[1]
            if re.match(r"\s*(Manufacturer|Type)\s+", line):
                columns = re.split(r"\s{2,}", line.strip())
                if len(columns) >= 4:
                    equipment[columns[0].lower()] = columns[1:1 + len(shaker_columns)]
            match = re.search(r"Shaker\s+([1-3])\s+(\d+(?:\.\d+)?)\s+\d+\s*$", line)
            if match and 0 <= float(match[2]) <= 24:
                runtimes[f"shaker_{match[1]}"] = float(match[2])
            match = re.search(r"MONGOOSE\s*(\d+)\s+([\d.\s]+)$", line)
            if match:
                values = match[2].split()
                # Final stock is the last column; daily use may be blank.
                if len(values) in (2, 3) and all(number(v) is not None for v in values):
                    screens[match[1]] = number(values[-1])
        requests = []
        for page in PdfReader(path).pages:
            fragments = []
            def visit(text, cm, tm, font, size):
                if text.strip():
                    fragments.append((tm[4], tm[5], unicodedata.normalize("NFKC", text).strip()))
            page.extract_text(visitor_text=visit)
            header = [y for x, y, value in fragments if "نیازها" in value]
            if header:
                rows = defaultdict(list)
                for x, y, value in fragments:
                    if 0 < y < max(header) and x < float(page.mediabox.width) / 2:
                        rows[round(y)].append((x, value))
                requests.extend(" ".join(t for x, t in sorted(rows[y], reverse=True)) for y in sorted(rows, reverse=True))
        installed = {key: " + ".join(f"{count}*{size}" for size, count in Counter(values).items()) + " API"
                     for key, values in decks.items()}
        date = re.search(r"\d{4}/\d{2}/\d{2}", text)
        return {"remarks": self._pdf_remarks(path, mud=True), "runtimes": runtimes, "screens": screens,
                "installed_screens": installed, "requests": requests, "equipment": equipment,
                "date": date[0] if date else None}

    def collect_source_details(self, files: dict, data: dict, logs: list[str]) -> dict:
        """Keep full source evidence outside the unchanged Tool Pusher workbook."""
        details = {"selected_reports": {k: p.name if p else None for k, p in files.items()},
                   "parsed_reports": data, "sources": [], "log": logs}
        for path in sorted(self.folder_path.iterdir()):
            if path.suffix.lower() not in {".xlsx", ".pdf"} or not path.is_file():
                continue
            if path.name.startswith("~$") or "consolidated" in path.stem.lower():
                continue
            if re.search(r"tool[ _-]*(?:poosher|pusher)", path.stem, re.I):
                continue
            source = {"report": path.name}
            details["sources"].append(source)
            try:
                if path.suffix.lower() == ".pdf":
                    source["pages"] = [{"page": n, "text": page.extract_text() or ""}
                                       for n, page in enumerate(PdfReader(path).pages, 1)]
                    if any(not p["text"].strip() for p in source["pages"]):
                        self._log(logs, f"Warning: {path.name} has pages requiring OCR.")
                else:
                    values = openpyxl.load_workbook(path, data_only=True)
                    formulas = None
                    try:
                        formulas = openpyxl.load_workbook(path, data_only=False)
                        source["sheets"] = []
                        for sheet in formulas:
                            if normalize(sheet.title) in {"activity codes", "setting"}:
                                continue
                            cells = []
                            for row in sheet:
                                for cell in row:
                                    if cell.value is None: continue
                                    entry = {"cell": cell.coordinate, "value": values[sheet.title][cell.coordinate].value}
                                    if cell.data_type == "f": entry["formula"] = cell.value
                                    cells.append(entry)
                            source["sheets"].append({"sheet": sheet.title, "cells": cells})
                    finally:
                        values.close()
                        if formulas: formulas.close()
            except Exception as exc:
                source["error"] = str(exc)
                self._log(logs, f"Warning: source details could not read {path.name}: {exc}")
        return details

    def run(self) -> tuple[bool, list[str]]:
        logs: list[str] = []
        if not self.folder_path.is_dir():
            self._log(logs, f"Error: '{self.folder_path}' is not a directory.")
            return False, logs
        try:
            files = self.locate_files()
            if files["target"]:
                return self._run_day(files, logs), logs
            days = [p for p in sorted(self.folder_path.iterdir()) if p.is_dir() and re.fullmatch(r"\d{4}[.-]\d{2}[.-]\d{2}", p.name)]
            if not days:
                raise ValueError("No Tool Pusher .xlsx target or dated daily subfolders found.")
            results = []
            for index, folder in enumerate(days, 1):
                self._set_status(f"Day {index}/{len(days)}: {folder.name}")
                self._log(logs, f"\n--- {folder.name} ---")
                child = Consolidator(folder, lambda msg: self._log(logs, msg), self.status_callback)
                success, _ = child.run()
                results.append(success)
                self.outputs.extend(child.outputs)
            self._log(logs, f"Batch finished: {sum(results)}/{len(results)} days saved. Review warnings in each day.")
            return all(results), logs
        except Exception as exc:
            self._log(logs, f"Error: {exc}")
            return False, logs

    def _run_day(self, files: dict[str, Path | None], logs: list[str]) -> bool:
        target = files["target"]
        self._log(logs, f"Target: {target.name}")
        report_date = re.search(r"\d{4}[./-]\d{2}[./-]\d{2}", target.stem)
        expected_date = re.sub(r"[.-]", "/", report_date[0]) if report_date else None
        data = {}
        for key, parser in (("pob", self.extract_pob), ("company_ddr", self.parse_company_man_ddr),
                            ("mdr", self.parse_mechanical_report), ("edr", self.parse_electrical_report),
                            ("mud_report", self.parse_mud_report)):
            path = files[key]
            if path is None:
                self._log(logs, f"Warning: {key.upper()} missing; existing target values retained.")
                continue
            self._set_status(f"Reading {path.name}...")
            try:
                data[key] = parser(path)
                self._log(logs, f"Read {key.upper()}: {path.name}")
                source_date = data[key].get("date")
                if expected_date and source_date and re.fullmatch(r"\d{4}/\d{2}/\d{2}", str(source_date)) and str(source_date) != expected_date:
                    self._log(logs, f"Warning: {path.name} contains date {source_date}; target is {expected_date}. Review this source's values.")
            except Exception as exc:
                self._log(logs, f"Warning: {key.upper()} could not be read: {exc}; section retained.")

        company_well = data.get("company_ddr", {}).get("fields", {}).get("well_name")
        expected_well = re.search(r"S[ _-]*(\d+)", str(company_well or ""), re.I)
        if expected_well:
            for source, report in data.items():
                if report.get("well_number") and report["well_number"] != expected_well[1]:
                    self._log(logs, f"Warning: {source.upper()} identifies well S-{report['well_number']}; company DDR identifies S-{expected_well[1]}. Review the source identity before issuing.")

        wb = openpyxl.load_workbook(target)
        temporary = None
        try:
            report = TargetReport(wb, lambda message: self._log(logs, message))
            report.company(data.get("company_ddr", {}))
            report.headers(expected_date, data.get('company_ddr', {}).get('fields', {}).get('rig_name'))
            report.personnel(data.get("pob", {}))
            report.runtimes(data)
            for source, label in (("mdr", "Mechanic"), ("edr", "Electrician")):
                report.remarks(label, data.get(source, {}).get("remarks", []))
                report.oil(data.get(source, {}))
            report.mud(data.get("mud_report", {}))
            self._set_status("Collecting source details...")
            details = self.collect_source_details(files, data, logs)
            details['target_values_from_sources'] = report.assignments
            details['section_results'] = report.section_results
            if not data:
                raise ValueError("No readable source reports were found; no output saved.")
            backup = target.with_suffix(".xlsx.bak")
            if not backup.exists():
                shutil.copy2(target, backup)
                self._log(logs, f"Backup: {backup.name}")
            out_file = target.with_name(f"{target.stem}_CONSOLIDATED.xlsx")
            index = 2
            while out_file.exists() or out_file.with_suffix(".details.json").exists():
                out_file = target.with_name(f"{target.stem}_CONSOLIDATED_{index}.xlsx")
                index += 1
            self._set_status("Saving report values...")
            with tempfile.NamedTemporaryFile(dir=target.parent, prefix=".consolidating-", suffix=".xlsx", delete=False) as handle:
                temporary = Path(handle.name)
            save_report_package(target, temporary, wb, report.changes)
            self._set_status('Reading back saved DDR values...')
            details['saved_value_check'] = check_saved_values(temporary, report.assignments)
            failures = [item for item in report.section_results if item['status'] == 'mapping_failed']
            if failures:
                self._log(logs, f'Warning: {len(failures)} mappings failed; these sections retain original values. See section_results in the companion file.')
            check = details['saved_value_check']
            self._log(logs, f"Saved value readback: {check['nonempty_values']} nonempty values and {check['assigned_cells'] - check['nonempty_values']} cleared slots match assignments. Retained template values are not counted as imports.")
            details_file = out_file.with_suffix(".details.json")
            self._log(logs, f"Saved: {out_file} ({report.count} values updated). Review warnings before issuing.")
            self._log(logs, f"Full source details and unmatched values: {details_file.name}")
            details_file.write_text(json.dumps(details, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
            temporary.rename(out_file)
            self.outputs.append(out_file)
            return True
        finally:
            wb.close()
            if temporary and temporary.exists():
                temporary.unlink()
