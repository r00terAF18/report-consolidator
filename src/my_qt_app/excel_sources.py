"""Read workbook reports using labels and table headers, independent of addresses."""
import re
from datetime import time

import openpyxl

from my_qt_app.report_layout import ReportLayout, LayoutError, label_key
from my_qt_app.operation_mapping import drilling_summary


def numeric(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (float, int)):
        return float(value)
    text = str(value or '').strip().replace(',', '').replace('٫', '.')
    if text == 'صفر':
        return 0.0
    return float(text) if re.fullmatch(r'\d+(?:\.\d+)?', text) else None


def read_company(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        # Sheet names also vary: identify the report pages by their labels.
        layouts = [ReportLayout(s) for s in wb]
        summary = next((s for s in layouts if s.matches('Rig Activity')), None)
        details = next((s for s in layouts if s.matches('Mud Type')), None)
        if not summary or not details:
            raise LayoutError('Company DDR pages could not be identified by their headings')
        fields, warnings, field_sources = {}, [], {}
        def read(key, layout, labels, *, latest=False):
            try:
                scope = {"start": layout.find('MUD Data').row, "end": layout.find('Volume data').row} if latest else {}
                label = layout.find(labels, **scope)
                cell = layout.right(label)
                if latest:
                    # Reading columns are defined by Report Time, not by letters.
                    header = layout.find('Report Time')
                    chemical = layout.find('Product Type')
                    columns = [c.column for c in layout.cells if c.row == header.row
                               and layout.bounds(header)[3] < c.column < chemical.column
                               and (isinstance(c.value, time)
                                    or isinstance(c.value, str) and re.fullmatch(r'\s*\d{1,2}:\d{2}(?::\d{2})?\s*', c.value))]
                    if not columns:
                        raise LayoutError('Mud reading time headers missing')
                    values = [layout.sheet.cell(label.row, col).value for col in columns]
                    populated = [(col, value) for col, value in zip(columns, values)
                                 if value is not None and str(value).strip()]
                    if populated:
                        column, value = populated[-1]
                        cell = layout.sheet.cell(label.row, column)
                    else:
                        value = None
                else:
                    value = cell.value
                if key == 'mud_type' and isinstance(value, str):
                    value = value.lstrip("'")
                fields[key] = value
                field_sources[key] = {'sheet': layout.sheet.title, 'cell': cell.coordinate,
                                      'label': label.value, 'value': value}
                return value
            except LayoutError as exc:
                warnings.append(str(exc))
                return None
        for key, labels in {
            'planned_td': 'Estimated Final Depth MD (m)', 'rig_days': 'Actual Rig Days',
            'report_no': 'Report No.', 'rig_name': 'Rig Name', 'field_name': 'Field', 'well_name': 'Well Name',
            'night_supervisor': 'Drilling Supervisor (Night)',
            'supervisor': 'Drilling Supervisor (Day)', 'geologist': 'Well Site Geologist',
            'previous_depth': 'MD (m)@ 0:00', 'current_depth': 'MD (m)@ 24:00',
            'morning_depth': 'MD (m)@ 6:00 AM', 'daily_footage': '24 hrs Drilled (m)',
            'daily_rop': 'AVG. ROP (m/hrs.)', 'kop': 'KOP #1',
        }.items():
            read(key, summary, labels)
        for key, labels in {
            'bit_no': 'Bit No.', 'bit_type': 'Bit Type', 'bit_size': 'Bit Size (inch)',
            'bit_serial': 'Bit Serial No.', 'iadc': 'IADC Code',
            'bit_drilled': 'Bit Cum. Drilled (m)', 'bit_hours': 'Bit Cum. Hrs on Bottom',
            'bit_condition': 'Bit Dull Condition', 'bit_rop': 'Cum.Bit Ave ROP (m/hr)', 'drilled_interval': 'Bit Drilled (m)',
            'drilling_hours': 'Bit Hours on Bottom', 'hole_volume': 'Vol. in Hole (bbls)',
            'bop_test': 'Last BOP Test', 'fire_drill': 'Last Fire Drill',
            'h2s_drill': 'Last H2S Drill', 'kick_drill': 'Last BOP Drill',
            'wind': 'Wind Speed (KPH)', 'temperature': 'TEMP(°C)', 'visibility': 'Visibility (km)',
        }.items():
            read(key, details, labels)
        for key, labels in {
            'mud_type': 'Mud Type', 'mud_weight': 'Mud Weight', 'viscosity': ('Funnel Viscosity', 'Vis'),
            'pv': 'PV (CP)', 'yp': 'YP', 'filtrate': 'API Fluid Loss (cc/30 min)',
            'ph': 'PH', 'chlorides': 'Chloride (mg/lit)', 'calcium': 'Hardness Ca++',
            'solids': 'Solid %', 'kcl': 'KCL %Wt/lb/bbl',
            'alkalinity': ('Alkalinity (pom) ml', 'Alkalinity (pm) ml'),
        }.items():
            read(key, details, labels, latest=True)
        if numeric(fields.get('solids')) is not None:
            fields['solids'] = numeric(fields['solids']) / 100
        kcl = re.fullmatch(r'\s*(\d+(?:\.\d+)?)\s*/\s*(\d+(?:\.\d+)?)\s*', str(fields.get('kcl') or ''))
        fields['kcl'] = round(float(kcl[2]) * 453592.37 / 158.987294928, 1) if kcl else None
        if numeric(fields.get('wind')) is not None:
            fields['wind'] = round(numeric(fields['wind']) / 1.852, 2)
        for key in ('kop', 'temperature', 'visibility'):
            fields[key] = numeric(fields.get(key))
        low, high = numeric(fields.get('previous_depth')), numeric(fields.get('current_depth'))
        if low is not None and high is not None:
            fields['daily_footage'] = high - low
        fields['daily_hours'] = fields.get('drilling_hours')
        for key in ('daily_rop', 'bit_rop'):
            if numeric(fields.get(key)) is not None:
                fields[key] = round(numeric(fields[key]), 2)
        if fields.get('bit_rop') is None and numeric(fields.get('drilling_hours')):
            footage = numeric(fields.get('drilled_interval'))
            if footage is not None:
                fields['bit_rop'] = round(footage / numeric(fields['drilling_hours']), 2)
        for key, label in {'wob': 'W.O.B (klb)', 'rpm': 'Surf. RPM ( rev/gal)',
                           'torque': 'Torque (kft-lbs)', 'spp': 'Pump Pressure (psi)'}.items():
            try:
                cell = details.right(details.find(label))
                values = (cell.value, details.right(cell).value)
                if all(v is not None for v in values):
                    fields[key] = ' - '.join(map(str, values))
            except LayoutError as exc:
                warnings.append(str(exc))
        try:
            header = details.find('Pump Output (gpm)')
            minimum = details.right(header)
            maximum = details.right(minimum)
            values = details.below(minimum).value, details.below(maximum).value
            if all(v is not None for v in values): fields['flow'] = ' - '.join(map(str, values))
        except LayoutError as exc: warnings.append(str(exc))
        try:
            losses = [numeric(details.value(label)) for label in ('Mud Lost Down Hole (bbls)', 'Mud Lost at Surface (bbls)')]
            if all(v is not None for v in losses): fields['mud_loss'] = sum(losses)
        except LayoutError as exc: warnings.append(str(exc))
        try:
            bulk = details.find('Bulk Data')
            for key, label in (('diesel', 'Fuel'), ('drill_water', 'D.W')):
                col = details.find(label, start=bulk.row, end=details.find('BOP Stack and Well Head').row).column
                values = {}
                for name in ('On Hand', 'Used', 'Received'):
                    row = details.find(name, start=bulk.row, end=details.find('BOP Stack and Well Head').row,
                                       column=bulk.column).row
                    values[name] = numeric(details.sheet.cell(row, col).value)
                for name, suffix in (('Used', 'used'), ('Received', 'received'), ('On Hand', 'stock')):
                    fields[f'{key}_{suffix}'] = values[name]
                if all(v is not None for v in values.values()):
                    fields[f'{key}_opening'] = values['On Hand'] + values['Used'] - values['Received']
        except LayoutError as exc: warnings.append(str(exc))
        inventory = []
        try:
            chemicals = details.find('Mud Chemical')
            product = details.find('Product Type')
            stock_col = details.find('On Hand', start=product.row, end=product.row + 1).column
            unit_col = details.find('Unit', start=product.row, end=product.row + 1).column
            stop = details.find('Cement and Cement Additives').row
            next_headers = details.matches('Serial Number', start=product.row + 1, column=product.column)
            if next_headers:
                stop = min(stop, min(c.row for c in next_headers))
            for product_cell in details.cells:
                if product.row < product_cell.row < stop and product_cell.column == product.column and product_cell.value:
                    inventory.append({'product': product_cell.value,
                                      'stock': details.sheet.cell(product_cell.row, stock_col).value,
                                      'unit': details.sheet.cell(product_cell.row, unit_col).value})
        except LayoutError as exc: warnings.append(str(exc))
        try:
            nozzle_label = details.find('Nozzle #1')
            nozzle_heading = details.right(nozzle_label)
            size_col = details.right(nozzle_heading).column
            nozzles = {}
            for label in details.matches('Nozzle #', prefix=True):
                qty = details.right(label).value
                size = details.sheet.cell(label.row, size_col).value
                if numeric(size) is not None:
                    count = numeric(qty) if numeric(qty) is not None else 1
                    nozzles[size] = nozzles.get(size, 0) + count
            if nozzles: fields['nozzles'] = '+'.join(f'{count:g}*{size}' for size, count in nozzles.items())
        except LayoutError as exc: warnings.append(str(exc))
        try:
            heading = summary.find('Previous Casing Information')
            size = summary.find('Size (in)', start=heading.row)
            depth_col = summary.find('To (m)', start=size.row, end=size.row + 1).column
            stop = min(c.row for c in summary.matches('Rig Activity'))
            from_col = summary.find('From (m)', start=size.row, end=size.row + 1).column
            casings = [(c.value, summary.sheet.cell(c.row, depth_col).value) for c in summary.cells
                       if size.row < c.row < stop and c.column == size.column and c.value is not None
                       and (summary.sheet.cell(c.row, from_col).value is not None
                            or summary.sheet.cell(c.row, depth_col).value is not None)]
            if casings:
                size_value, depth = casings[-1]
                if depth is None:
                    warnings.append('Casing rows after the last complete entry have no shoe depth; target casing retained rather than substituting an earlier string.')
                else:
                    fields['casing'], fields['casing_shoe'] = f'Casing {size_value}@{depth}M', depth
        except LayoutError as exc: warnings.append(str(exc))
        try:
            label = summary.find('Formation')
            boundary = summary.find('Previous Casing Information').column
            values = [c for c in summary.cells if c.row == label.row
                      and summary.bounds(label)[3] < c.column < boundary
                      and c.value is not None and str(c.value).strip()]
            if len(values) != 1:
                raise LayoutError('Formation value missing or ambiguous before casing section')
            formation = str(values[0].value).split('@')[0].strip()
            if formation:
                fields['formation'] = formation
                heading = details.find('Formation Top Data')
                name = details.find('Name', start=heading.row, column=heading.column)
                row = details.find(formation, start=name.row, column=name.column).row
                for key, label in (('formation_top', 'm MD'), ('lithology', 'Lithology')):
                    col = details.find(label, start=name.row, end=name.row + 1).column
                    fields[key] = details.sheet.cell(row, col).value
        except LayoutError as exc: warnings.append(str(exc))
        try:
            heading = details.find('SURVEY')
            header = details.find('MD (m)', start=heading.row, column=heading.column)
            rows = details.region_rows(header, ('Weather Data', 'ENVIRONMENTAL SOLUTION'), column=None)
            for key, label in (('survey_md', 'MD (m)'), ('survey_tvd', 'TVD (m)'), ('survey_inc', 'Inc.'), ('survey_az', 'Azi')):
                col = details.find(label, start=header.row, end=header.row + 1).column
                values = [details.sheet.cell(r, col).value for r in rows if numeric(details.sheet.cell(r, header.column).value) is not None
                          and details.sheet.cell(r, col).value is not None]
                if values: fields[key] = ' / '.join(map(str, values))
        except LayoutError as exc: warnings.append(str(exc))
        bha = []
        for heading in details.matches('BHA No.', prefix=True):
            try:
                match = re.search(r'=\s*(\d+)', str(heading.value))
                if match: fields['bha_no'] = f'#{match[1]}'
                item = details.find('Item', start=heading.row, column=heading.column)
                headers = [c for c in details.cells if c.row == item.row and c.column >= item.column]
                # Stop at the next Item header; the two BHA tables have the same headings.
                next_items = [c.column for c in headers if c.column > item.column and label_key(c.value) == label_key('Item')]
                limit = min(next_items) if next_items else details.find('MUD Data').column
                cols = {name: next(c.column for c in headers if c.column < limit and label_key(c.value) == label_key(label))
                        for name, label in (('od', 'OD (in)'), ('length', 'Length (m)'))}
                cumulative_col = next((c.column for c in headers if c.column < limit and label_key(c.value) == label_key('Cum. Len (m)')), None)
                last_cumulative = None
                table_lengths = []
                for r in details.region_rows(item, 'TOTAL(m)', column=item.column):
                    text = details.sheet.cell(r, item.column).value
                    if not text: continue
                    qty = re.match(r'\s*(\d+)\s+(?:jts?|stds?|stands?)\s+(.+)', str(text), re.I)
                    name = qty[2] if qty else text
                    count = int(qty[1]) if qty else 1
                    if qty and re.search(r'\b(?:stds?|stands?)\b', str(text), re.I):
                        count *= 3  # The reference BHA lists two stands as six joints.
                    bha.append((name, details.sheet.cell(r, cols['od']).value, count,
                                details.sheet.cell(r, cols['length']).value))
                    table_lengths.append(numeric(details.sheet.cell(r, cols['length']).value))
                    if cumulative_col is not None:
                        last_cumulative = numeric(details.sheet.cell(r, cumulative_col).value)
                if table_lengths and all(v is not None for v in table_lengths) and last_cumulative is not None:
                    length_sum = round(sum(table_lengths), 3)
                    if abs(length_sum - last_cumulative) > 0.001:
                        warnings.append(f'BHA item lengths sum to {length_sum:g} m, but the source cumulative length is {last_cumulative:g} m; item lengths retained.')
            except (LayoutError, StopIteration) as exc: warnings.append(f'BHA: {exc}')
        try:
            heading = details.find('BOP Stack and Well Head')
            header = details.find('Name', start=heading.row, column=heading.column)
            columns = [details.find(label, start=header.row, end=header.row + 1).column for label in ('TYPE', 'SIZE', 'W.P (psi)', 'RAMS')]
            for key, label in (('annular_bop', 'HYDRIL'), ('single_bop', "TOP RAM'S"), ('bottom_bop', 'DOUBLE(BLIND/BTM)')):
                item = details.find(label, start=header.row, column=header.column)
                values = [details.sheet.cell(item.row, c).value for c in columns]
                fields[key] = ' '.join(str(v) for v in values if v is not None)
                if key == 'bottom_bop':
                    rams = re.fullmatch(r'DOUBLE\s*\(\s*(.*?)\s*/\s*(.*?)\s*\)', str(values[-1] or ''), re.I)
                    if rams:
                        base = ' '.join(str(v) for v in values[:-1] if v is not None)
                        fields['shear_bop'], fields['bottom_bop'] = f'{base} {rams[1]}', f'{base} {rams[2]}'
            casing_head = details.find('C.H.H', start=header.row, column=header.column)
            fields['bop_size'] = ' '.join(str(details.sheet.cell(casing_head.row, c).value)
                                           for c in columns[1:3] if details.sheet.cell(casing_head.row, c).value is not None)
        except LayoutError as exc: warnings.append(str(exc))
        pump_liners = {}
        try:
            liner = str(details.value('Pump Liner Size (inch)') or '')
            match = re.search(r'Pump\s*#?\s*([1-3](?:\s*[,/&]\s*[1-3])*)\s*:\s*([^xX×]+)[xX×]', liner, re.I)
            if match:
                pump_liners = {n: match[2].strip() for n in re.findall(r'[1-3]', match[1])}
            elif liner:
                warnings.append(f'Pump liner specification has no explicit pump/liner assignment: {liner}')
        except LayoutError as exc: warnings.append(str(exc))
        scr = []
        headings = details.matches('SCR Data', prefix=True)
        if len(headings) == 1:
            start = headings[0]
            stop = min(c.row for c in details.matches('BHA No.', start=start.row, prefix=True))
            for label in details.matches('Pump No', start=start.row, end=stop, prefix=True):
                pump = re.search(r'([1-3])', label.value)
                spm_cell = details.right(label)
                unit = details.right(spm_cell)
                if pump and label_key(unit.value) == 'stk' and numeric(spm_cell.value) is not None and details.right(unit).value is not None:
                    scr.append((int(pump[1]), spm_cell.value, details.right(unit).value))
        shakers = []
        try:
            heading = details.find('Solid Control')
            header = details.find('Equipment', start=heading.row, column=heading.column)
            screen_col = details.find('Size/# Cones', start=header.row, end=header.row + 1).column
            hours_col = details.find('Daily Hrs/ Cum hrs', start=header.row, end=header.row + 1).column
            for label in details.matches('Shaker #', start=header.row, column=header.column, prefix=True):
                n = int(re.search(r'\d+', label.value)[0])
                shakers.append((n, details.sheet.cell(label.row, screen_col).value,
                                numeric(str(details.sheet.cell(label.row, hours_col).value or '').split('/')[0])))
        except LayoutError as exc: warnings.append(str(exc))
        operations = []
        for heading in summary.matches('Rig Activity'):
            records = []
            try:
                row = heading.row
                headers = [summary.find(label, start=row, end=row + 1) for label in ('From', 'To', 'hrs', 'Code', 'Sub Code')]
                totals = summary.matches('Total', start=row + 1, column=headers[0].column)
                if not totals:
                    raise LayoutError('Activity table Total boundary missing')
                stop = min(c.row for c in totals)
                for r in range(summary.bounds(heading)[2] + 1, stop):
                    description = summary.sheet.cell(r, heading.column).value
                    if not description or not str(description).strip(): continue
                    values = [summary.sheet.cell(r, h.column).value for h in headers]
                    begin, end, hours = values[:3]
                    if isinstance(begin, time) and isinstance(end, time):
                        duration = ((end.hour - begin.hour) * 3600 + (end.minute - begin.minute) * 60 + end.second - begin.second) % 86400 / 3600
                        values[2] = duration or numeric(hours) or 24
                    records.append((*values, description))
            except LayoutError as exc:
                warnings.append(f'Activity table at {summary.sheet.title}!{heading.coordinate}: {exc}')
            operations.append(records)
        categories = {}
        try:
            breakdown = summary.find('Time Break Down')
            code = summary.find('no', start=breakdown.row, column=breakdown.column)
            name_col = summary.right(code).column
            totals = summary.matches('Total', start=code.row + 1, column=name_col)
            if not totals:
                raise LayoutError('Time Break Down Total boundary missing')
            stop = min(c.row for c in totals)
            for row in range(code.row + 1, stop):
                value = summary.sheet.cell(row, code.column).value
                if numeric(value) is not None:
                    categories[str(int(value))] = str(summary.sheet.cell(row, name_col).value or '')
        except LayoutError as exc: warnings.append(str(exc))
        daily = operations[0] if operations else []
        if not fields.get('bit_condition') and len(operations) > 1:
            conditions = [match[1] for record in operations[1]
                          if re.search(r'\bL\s*/\s*D\b', str(record[5]), re.I)
                          for match in re.finditer(r'\bDULL\s+COND(?:ITION)?\s*\(([^)]+)\)', str(record[5]), re.I)]
            if len(conditions) == 1:
                fields['bit_condition'] = conditions[0]
        drilling = drilling_summary(daily, categories, numeric(fields.get('drilled_interval')))
        if drilling['rotating_hours']:
            fields['drilling_hours'] = fields['daily_hours'] = drilling['rotating_hours']
        if drilling['drill_out_hours']:
            fields['daily_hours'] = (f"{drilling['rotating_hours']:g} "
                                     f"({drilling['drill_out_hours']:g} D/O + {drilling['drilling_hours']:g} DRLG)")
            if 'drill_out_footage' in drilling:
                fields['daily_footage'] = f"{drilling['drill_out_footage']:g} m D/O + {drilling['new_hole_footage']:g} m DRLG"
        reported_footage = field_sources.get('daily_footage', {}).get('value')
        if 'new_hole_footage' in drilling and numeric(reported_footage) != drilling['new_hole_footage']:
            warnings.append(f"Company reports {reported_footage} m daily footage, but written new-hole intervals total {drilling['new_hole_footage']:g} m; review the discrepancy.")
        extra = {}
        for key, label in (('summary', 'Summary of Activities in Last 24:00 hours'),
                           ('forecast', 'Operation Forecast for next 24 hours'), ('note', 'Note')):
            try: extra[key] = summary.below(summary.find(label)).value
            except LayoutError as exc: warnings.append(str(exc))
        for key, label, stop_label in (('requests', 'Out Standing', 'Back Load'), ('received', 'Received', None)):
            try:
                heading = summary.find(label)
                col = summary.right(heading).column
                if stop_label:
                    rows = range(heading.row, summary.find(stop_label, start=heading.row, column=heading.column).row)
                else:
                    # Receipts occupy the last section of the source form.
                    rows = range(heading.row, summary.sheet.max_row + 1)
                extra[key] = [str(summary.sheet.cell(r, col).value) for r in rows if summary.sheet.cell(r, col).value]
            except LayoutError as exc: warnings.append(str(exc))
        return {'fields': fields, 'field_sources': field_sources, 'bha': bha, 'pump_liners': pump_liners, 'scr': scr, 'shakers': shakers,
                'inventory': inventory, 'operation_categories': categories, 'drilling_summary': drilling,
                'operations': operations[0] if operations else [], 'morning': operations[1] if len(operations) > 1 else [],
                'warnings': warnings, **extra}
    finally:
        wb.close()


def read_service(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        s = ReportLayout(wb.active)
        heading = s.find('شرح اقدامات انجام شده')
        rows = s.region_rows(heading, 'قطعات مصرف شده', column=heading.column)
        remarks = []
        for row in rows:
            value = s.sheet.cell(row, heading.column).value
            if value:
                remarks.extend(t.strip() for t in re.split(r'(?:^|\n)\s*\d+\s*[-_]\s*', str(value)) if t.strip())
        runtimes = {}
        for cell in s.cells:
            key = label_key(cell.value)
            motor = re.fullmatch(r'موتور([1-4])', key)
            compressor = re.fullmatch(r'کمپرسورشماره([1-3])', key)
            name = f'gen_{motor[1]}' if motor else f'comp_{compressor[1]}' if compressor else 'emergency' if key == 'اضطراری' else None
            if name:
                value = numeric(s.below(cell).value)
                if value is not None: runtimes[name] = value
        def table(section, columns, stop):
            title = s.find(section)
            first = s.find('ردیف', start=title.row, end=s.find(stop).row, column=title.column)
            cols = {name: s.find(label, start=first.row, end=first.row + 1).column for name, label in columns.items()}
            return [{name: s.sheet.cell(r, col).value for name, col in cols.items()} for r in s.region_rows(first, stop, column=title.column)
                    if s.sheet.cell(r, cols['description']).value]
        materials = table('قطعات مصرف شده', {'code': 'کد فنی کالا یا MESC', 'description': 'شرح کالا', 'unit': 'واحد', 'quantity': 'تعداد'}, 'مقدار مصرف روغن و مواد مصرفی')
        oils = table('مقدار مصرف روغن و مواد مصرفی', {'description': 'شرح قلم', 'unit': 'واحد', 'used': 'مقدار مصرف'}, 'ساعت کارکرد تجهیزات')
        return {'remarks': remarks, 'runtimes': runtimes, 'materials': materials, 'oils': oils, 'date': s.value('تاریخ')}
    finally:
        wb.close()


def read_pob(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        s = ReportLayout(wb.active)
        headers = s.matches('Name')
        if len(headers) != 2: raise LayoutError('POB name headers missing or ambiguous')
        group_aliases = {'dpr': ('DPR.',), 'dpr_services': ('DPR Services',), 'client': ('Client',),
                         'services': ('Client.Services',), 'catering': ('CATERING',), 'guard': ('Guard',), 'labour': ('Labour',)}
        groups, people, staff = {}, [], {}
        total = s.find('Total POB')
        captions = {}
        for name, aliases in group_aliases.items():
            candidates = s.matches(aliases, end=total.row)
            if len(candidates) == 1:
                heading = candidates[0]
                captions[name] = (heading.row, heading.column, s.bounds(heading)[2])
            elif name != 'guard':
                raise LayoutError(f'POB group missing or ambiguous: {name}')
        if 'guard' not in captions:
            # Some reports omit the Guard caption. Recover that group from its
            # explicit guard designations in the right personnel table.
            names = max(headers, key=lambda c: c.column)
            designation = s.find('Designation', start=names.row, end=names.row + 1,
                                 column=s.right(names).column)
            guard_rows = [r for r in range(captions['catering'][0] + 1, captions['labour'][0])
                          if s.sheet.cell(r, names.column).value
                          and 'حراست' in label_key(s.sheet.cell(r, designation.column).value)]
            if not guard_rows: raise LayoutError('POB Guard heading and designations missing')
            captions['guard'] = (min(guard_rows) - 1, captions['labour'][1], min(guard_rows) - 1)
        for name, (row, column, end_row) in captions.items():
            names = min((h for h in headers if h.column > column), key=lambda h: h.column)
            lane_start = max((h.column for h in headers if h.column < names.column), default=0)
            stops = [r for r, c, _ in captions.values() if r > row and lane_start < c < names.column]
            stop = min(stops + [total.row])
            designation = s.find('Designation', start=names.row, end=names.row + 1,
                                 column=s.right(names).column)
            entries = [(s.sheet.cell(r, names.column).value, s.sheet.cell(r, designation.column).value)
                       for r in range(end_row + 1, stop) if s.sheet.cell(r, names.column).value]
            groups[name] = len(entries)
            if name == 'dpr': people = entries
            for key, title in {'senior_tool_pusher': 'رئیس دستگاه', 'junior_tool_pusher': 'سرحفار',
                               'top_client_rep': 'ناظر عالیه کارفرما'}.items():
                matches = [person for person, designation in entries if label_key(designation) == label_key(title)]
                if len(matches) == 1:
                    staff[key] = str(matches[0]).strip()
        roles = {key: 0 for key in ('tool_pusher', 'driller', 'assistant_driller', 'mud_engineer', 'electrical', 'mechanical', 'trainee')}
        for _, title in people:
            key = label_key(title)
            for role, names in {'tool_pusher': ('رئیس دستگاه', 'سرحفار'), 'driller': ('حفار',),
                                'assistant_driller': ('کمک حفار',), 'mud_engineer': ('گلشناس روز', 'گلشناس شبکار', 'گلشناس شب'),
                                'electrical': ('برقکارارشد', 'برقکار2'), 'mechanical': ('مکانیک ارشد', 'مکانیک2'),
                                'trainee': ('کارآموز', 'trainee')}.items():
                if key in map(label_key, names): roles[role] += 1; break
        roles.update(rig_crew=groups['dpr'] + groups['dpr_services'] - sum(roles.values()), casual_lab=groups['labour'],
                     catering=groups['catering'], services=groups['services'], client=groups['client'], security=groups['guard'])
        dates = [str(c.value) for c in s.cells if c.row < min(h.row for h in headers) and re.fullmatch(r'\d{4}/\d{2}/\d{2}', str(c.value or ''))]
        return {'fields': roles, 'staff': staff, 'groups': groups, 'total': sum(groups.values()),
                'date': dates[0] if len(dates) == 1 else None}
    finally:
        wb.close()
